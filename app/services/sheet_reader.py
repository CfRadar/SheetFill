"""Sheet reader service for reading .xls, .xlsm, and .xlsx Flipkart listings."""

import os
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
import openpyxl
from openpyxl.utils import get_column_letter
import xlrd
from pydantic import BaseModel
import zipfile


class RawRow(BaseModel):
    """Represents an unparsed raw data row from the product sheet."""

    row_index: int  # 1-based index (5..500)
    cells: Dict[str, Any]  # Col letter (A..CA) -> str | float | int | None

    def get(self, col: str, default: Any = None) -> Any:
        return self.cells.get(col, default)

    def is_empty_in_range(self, start_col: str = "G", end_col: str = "BP", ignore_cols: Optional[List[str]] = None) -> bool:
        """Checks if all cells in the specified range are empty, ignoring specific columns."""
        ignore = set(ignore_cols or [])
        start_idx = openpyxl.utils.column_index_from_string(start_col)
        end_idx = openpyxl.utils.column_index_from_string(end_col)
        for idx in range(start_idx, end_idx + 1):
            letter = get_column_letter(idx)
            if letter in ignore:
                continue
            val = self.cells.get(letter)
            if val is not None and str(val).strip():
                return False
        return True


class RawSheet(BaseModel):
    """Represents the raw sheet data and headers."""

    headers: Dict[str, str]  # Col letter -> Header name from Row 1
    rows: List[RawRow]
    total_rows_read: int
    sheet_name: str = "t_shirt"
    product_type: str = "t_shirt"


def _clean_cell_value(val: Any) -> Any:
    """Normalize cell values: trim text, convert floats that are ints, treat whitespace as None."""
    if val is None:
        return None
    if isinstance(val, str):
        trimmed = val.strip()
        return trimmed if trimmed else None
    if isinstance(val, (int,)):
        return val
    if isinstance(val, float):
        if val.is_integer():
            return int(val)
        return val
    s = str(val).strip()
    return s if s else None


SUPPORTED_PRODUCT_SHEETS = {
    "t_shirt": "t_shirt",
    "tshirt": "t_shirt",
    "sweatshirt": "sweatshirt",
}


def detect_product_sheet(sheet_names: List[str]) -> Tuple[str, str]:
    """Finds the product worksheet (t_shirt or sweatshirt) from a list of sheet names."""
    for name in sheet_names:
        clean = name.lower().strip().replace(" ", "").replace("-", "_")
        if clean in SUPPORTED_PRODUCT_SHEETS:
            return name, SUPPORTED_PRODUCT_SHEETS[clean]
    raise ValueError(
        f"Neither 't_shirt' nor 'sweatshirt' sheet found in uploaded workbook. Available sheets: {sheet_names}"
    )


def get_master_headers(master_path: Optional[str] = None, product_type: str = "t_shirt") -> Dict[str, str]:
    """Reads header row 1 from the master template (.xlsm or .xls)."""
    base_dir = Path(__file__).resolve().parent.parent.parent
    if not master_path:
        if product_type == "sweatshirt":
            master_path = str(base_dir / "templates" / "Flipkart_Sweatshirt_Template.xls")
        else:
            master_path = str(base_dir / "templates" / "Flipkart_TShirt_Template.xlsm")

    m_path = Path(master_path)
    if not m_path.exists():
        # Fallback to .xls if .xlsm not found or vice versa
        if m_path.suffix.lower() == ".xlsm" and m_path.with_suffix(".xls").exists():
            m_path = m_path.with_suffix(".xls")
        elif m_path.suffix.lower() == ".xls" and m_path.with_suffix(".xlsm").exists():
            m_path = m_path.with_suffix(".xlsm")

    target_sheet = "sweatshirt" if product_type == "sweatshirt" else "t_shirt"

    if m_path.suffix.lower() == ".xls" and not zipfile.is_zipfile(str(m_path)):
        wb_xls = xlrd.open_workbook(str(m_path), formatting_info=False)
        actual_sheet, _ = detect_product_sheet(wb_xls.sheet_names())
        sh = wb_xls.sheet_by_name(actual_sheet)
        headers = {}
        for c in range(sh.ncols):
            col_letter = get_column_letter(c + 1)
            hdr_val = sh.cell_value(0, c)
            headers[col_letter] = str(hdr_val).strip() if hdr_val is not None else ""
        return headers

    wb = openpyxl.load_workbook(str(m_path), read_only=True, data_only=True)
    try:
        actual_sheet, _ = detect_product_sheet(wb.sheetnames)
        sh = wb[actual_sheet]
        headers = {}
        max_col = sh.max_column or (79 if product_type == "sweatshirt" else 68)
        for c in range(1, max_col + 1):
            col_letter = get_column_letter(c)
            hdr_val = sh.cell(1, c).value
            headers[col_letter] = str(hdr_val).strip() if hdr_val is not None else ""
        return headers
    finally:
        wb.close()


def read_tshirt_sheet(
    file_path: str,
    master_template_path: Optional[str] = None,
    product_type: Optional[str] = None,
) -> RawSheet:
    """Reads product sheet ('t_shirt' or 'sweatshirt') from an uploaded .xls, .xlsm, or .xlsx file.
    
    Validates:
    - Product sheet exists.
    - Row 1 headers match master template headers for that product category.
    
    Reads data rows 5..500.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = path.suffix.lower()
    if ext not in [".xls", ".xlsm", ".xlsx"]:
        raise ValueError(f"Unsupported file format '{ext}'. Must be .xls, .xlsm, or .xlsx")

    headers: Dict[str, str] = {}
    raw_rows: List[RawRow] = []

    if ext == ".xls" and not zipfile.is_zipfile(file_path):
        wb_xls = xlrd.open_workbook(file_path, formatting_info=False)
        sheet_name, detected_type = detect_product_sheet(wb_xls.sheet_names())
        final_type = product_type or detected_type
        sh = wb_xls.sheet_by_name(sheet_name)

        master_headers = get_master_headers(master_template_path, product_type=final_type)

        # Validate header row 1 (0-indexed in xlrd: row 0)
        target_cols = max(sh.ncols, len(master_headers))
        for c in range(sh.ncols):
            col_letter = get_column_letter(c + 1)
            val = sh.cell_value(0, c)
            headers[col_letter] = str(val).strip() if val is not None else ""

        # Check header alignment
        _validate_headers(headers, master_headers)

        # Read rows 5..500 (0-indexed 4..499)
        num_cols = len(master_headers)
        for r_idx in range(4, min(500, sh.nrows)):
            cells = {}
            for c_idx in range(sh.ncols):
                col_letter = get_column_letter(c_idx + 1)
                raw_val = sh.cell_value(r_idx, c_idx)
                cells[col_letter] = _clean_cell_value(raw_val)
            # Fill remaining columns up to num_cols if ncols < num_cols
            for c_idx in range(sh.ncols, num_cols):
                cells[get_column_letter(c_idx + 1)] = None
            raw_rows.append(RawRow(row_index=r_idx + 1, cells=cells))

    else:
        # If the file ends with .xls but is a zip archive, pass a file-like object so openpyxl doesn't reject .xls
        f_target = open(file_path, "rb") if ext == ".xls" else file_path
        try:
            wb = openpyxl.load_workbook(f_target, read_only=False, keep_vba=False, data_only=True)
            sheet_name, detected_type = detect_product_sheet(wb.sheetnames)
            final_type = product_type or detected_type
            sh = wb[sheet_name]

            master_headers = get_master_headers(master_template_path, product_type=final_type)
            num_cols = len(master_headers)

            # Validate header row 1
            for c in range(1, num_cols + 1):
                col_letter = get_column_letter(c)
                val = sh.cell(1, c).value
                headers[col_letter] = str(val).strip() if val is not None else ""

            _validate_headers(headers, master_headers)

            # Read rows 5..500
            max_r = min(500, max(sh.max_row or 5, 5))
            for r in range(5, max_r + 1):
                cells = {}
                for c in range(1, num_cols + 1):
                    col_letter = get_column_letter(c)
                    raw_val = sh.cell(r, c).value
                    cells[col_letter] = _clean_cell_value(raw_val)
                raw_rows.append(RawRow(row_index=r, cells=cells))
        finally:
            if hasattr(f_target, "close"):
                f_target.close()

    return RawSheet(
        headers=headers,
        rows=raw_rows,
        total_rows_read=len(raw_rows),
        sheet_name=sheet_name,
        product_type=final_type,
    )


def _validate_headers(headers: Dict[str, str], master_headers: Dict[str, str]) -> None:
    """Validates that uploaded headers match master template headers."""
    mismatches = []
    for col_letter, expected in master_headers.items():
        actual = headers.get(col_letter, "")
        # Compare normalized
        if expected.lower() != actual.lower():
            mismatches.append(f"Col {col_letter}: expected '{expected}', found '{actual}'")

    if mismatches:
        preview = "; ".join(mismatches[:5])
        raise ValueError(
            f"Uploaded workbook row-1 headers do not match Flipkart master template ({len(mismatches)} mismatches: {preview})"
        )
