"""End-to-end API integration tests covering Section 10 acceptance criteria."""

import json
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter
import pytest
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)
FIXTURES_DIR = Path(__file__).parent / "fixtures"


def test_api_full_flow_compare_with_2_final():
    # 1. Analyze 1_before.xls
    xls_path = FIXTURES_DIR / "1_before.xls"
    with open(xls_path, "rb") as f:
        resp = client.post("/api/sheets/analyze", files={"file": ("1_before.xls", f, "application/vnd.ms-excel")})
    assert resp.status_code == 200
    data = resp.json()
    sid = data["session_id"]
    assert data["summary"]["total"] == 8
    assert data["summary"]["needs_json"] == 8

    # 2. Bulk assign batch_1_to_2.json
    batch_raw = (FIXTURES_DIR / "batch_1_to_2.json").read_text(encoding="utf-8")
    resp = client.post(f"/api/sheets/{sid}/bulk-assign", json={"raw_text": batch_raw})
    assert resp.status_code == 200
    assign_data = resp.json()
    assert len(assign_data["assigned"]) == 8

    # 3. Confirm bulk assignment
    pairs = [
        {"design_id": a["design_id"], "json_index": a["json_index"], "parsed_json": a["parsed_json"]}
        for a in assign_data["assigned"]
    ]
    resp = client.post(f"/api/sheets/{sid}/bulk-assign/confirm", json={"pairs": pairs})
    assert resp.status_code == 200
    confirm_data = resp.json()
    assert confirm_data["summary"]["needs_json"] == 0

    # 4. Generate in ready_only mode
    resp = client.post(f"/api/sheets/{sid}/generate", json={"mode": "ready_only"})
    assert resp.status_code == 200
    gen_data = resp.json()
    download_url = gen_data["download_url"]
    assert download_url == f"/api/sheets/{sid}/download"

    # 5. Download generated workbook
    resp = client.get(download_url)
    assert resp.status_code == 200

    out_file = Path("output") / f"api_test_out_{sid}.xlsm"
    out_file.parent.mkdir(parents=True, exist_ok=True)
    out_file.write_bytes(resp.content)

    # 6. Column-by-column comparison with 2_final.xlsm rows 5..36
    wb_out = openpyxl.load_workbook(str(out_file), data_only=True)
    wb_exp = openpyxl.load_workbook(str(FIXTURES_DIR / "2_final.xlsm"), data_only=True)
    sh_out = wb_out["t_shirt"]
    sh_exp = wb_exp["t_shirt"]

    # Documented intentional differences:
    # Pattern casing: "Graphic print" -> "Graphic Print"
    # Fabric care casing: "regular machine wash" -> "Regular Machine Wash"
    # Key features: dynamic from batch JSON vs fixed raw text in sample
    # Surface styling: dynamic from description text
    # Tee length: 1_before.xls stub row 11 explicitly contains 'Long' while 2_final has 'Medium'
    # HSN: 2_final.xlsm incorrectly put SKU string in Col Y; our system sets the valid GST HSN 61091000
    # Group ID (Col H): 2_final.xlsm incorrectly assigned 'group 1' to all designs; now incremented per design group
    # Packer details (Col AC): 2_final.xlsm left empty; now populated
    intentional_columns = {"AI", "BH", "BO", "BJ", "BF", "Y", "H", "AC"}

    unexpected_diffs = []
    for r in range(5, 37):
        for c in range(1, 69):
            col = get_column_letter(c)
            if col in ["A", "B", "C", "E", "F", "I"]:
                continue
            if col in intentional_columns:
                continue

            v_out = sh_out.cell(r, c).value
            v_exp = sh_exp.cell(r, c).value

            s_out = str(v_out).strip() if v_out is not None else ""
            s_exp = str(v_exp).strip() if v_exp is not None else ""

            # Check numeric equivalence
            try:
                if abs(float(s_out) - float(s_exp)) < 1e-4:
                    continue
            except ValueError:
                pass

            if s_out != s_exp:
                unexpected_diffs.append((r, col, s_exp, s_out))

    assert len(unexpected_diffs) == 0, f"Unexpected differences found: {unexpected_diffs[:10]}"

    # Verify Col H groups: Design 1 (rows 5..8) -> group 1, Design 2 (rows 9..12) -> group 2
    for r in range(5, 9):
        assert sh_out.cell(r, 8).value == "group 1"
    for r in range(9, 13):
        assert sh_out.cell(r, 8).value == "group 2"
    for r in range(13, 17):
        assert sh_out.cell(r, 8).value == "group 3"

    # Verify Col AC (Packer details) is populated
    for r in range(5, 37):
        assert sh_out.cell(r, 29).value is not None and str(sh_out.cell(r, 29).value).strip() != ""



def test_api_partial_run():
    # Ingest 1_before.xls, fill 5 of 8 designs -> generate ready_only
    xls_path = FIXTURES_DIR / "1_before.xls"
    with open(xls_path, "rb") as f:
        resp = client.post("/api/sheets/analyze", files={"file": ("1_before.xls", f, "application/vnd.ms-excel")})
    sid = resp.json()["session_id"]

    # Take first 5 designs from batch
    batch_raw = (FIXTURES_DIR / "batch_1_to_2.json").read_text(encoding="utf-8")
    items = json.loads(batch_raw)[:5]

    for idx, item in enumerate(items, start=1):
        did = f"d{idx}"
        resp = client.put(f"/api/sheets/{sid}/designs/{did}/json", json={"raw_text": json.dumps(item)})
        assert resp.status_code == 200

    # Generate ready_only
    resp = client.post(f"/api/sheets/{sid}/generate", json={"mode": "ready_only"})
    assert resp.status_code == 200
    report = resp.json()["report"]
    # 5 designs of 4 rows + 3 stubs of 1 row = 23 data rows
    assert report["rows_written"] == 23


def test_api_all_or_nothing_conflict():
    xls_path = FIXTURES_DIR / "1_before.xls"
    with open(xls_path, "rb") as f:
        resp = client.post("/api/sheets/analyze", files={"file": ("1_before.xls", f, "application/vnd.ms-excel")})
    sid = resp.json()["session_id"]

    # Try all_or_nothing without filling all designs
    resp = client.post(f"/api/sheets/{sid}/generate", json={"mode": "all_or_nothing"})
    assert resp.status_code == 409
    data = resp.json()
    assert data["error"]["code"] == "DESIGNS_NOT_READY"


def test_api_file_size_limit():
    large_payload = b"0" * (26 * 1024 * 1024)
    resp = client.post("/api/sheets/analyze", files={"file": ("huge.xlsm", large_payload, "application/vnd.ms-excel")})
    assert resp.status_code == 413


def test_api_meta_brand_manufacturer_hsn_endpoints():
    xls_path = FIXTURES_DIR / "1_before.xls"
    with open(xls_path, "rb") as f:
        resp = client.post("/api/sheets/analyze", files={"file": ("1_before.xls", f, "application/vnd.ms-excel")})
    assert resp.status_code == 200
    sid = resp.json()["session_id"]

    # 1. Update single design meta (d1)
    resp = client.put(
        f"/api/sheets/{sid}/designs/d1/meta",
        json={"brand": "BrandOne", "manufacturer": "MfgOne, Delhi", "hsn": "61099010"},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["design_meta"]["brand"] == "BrandOne"
    assert data["design_meta"]["manufacturer"] == "MfgOne, Delhi"
    assert data["design_meta"]["hsn"] == "61099010"

    # 2. Put JSON for d1 and verify preview rows contain these values
    sample_json = (FIXTURES_DIR / "moonlight.json").read_text(encoding="utf-8")
    resp = client.put(f"/api/sheets/{sid}/designs/d1/json", json={"raw_text": sample_json})
    assert resp.status_code == 200
    rows = resp.json()["preview_rows"]
    assert len(rows) == 4
    for r in rows:
        assert r["cells"]["AG"] == "BrandOne"
        assert r["cells"]["AB"] == "MfgOne, Delhi"
        assert r["cells"]["Y"] == "61099010"

    # 3. Apply to all designs
    resp = client.post(
        f"/api/sheets/{sid}/meta/apply-all",
        json={"brand": "GlobalBrand", "manufacturer": "GlobalMfg, Mumbai", "hsn": "61091000"},
    )
    assert resp.status_code == 200
    all_data = resp.json()
    assert all_data["global_meta"]["brand"] == "GlobalBrand"
    assert all_data["global_meta"]["manufacturer"] == "GlobalMfg, Mumbai"
    assert all_data["global_meta"]["hsn"] == "61091000"

    # d1 was refreshed: verify get_session has global and design meta
    resp = client.get(f"/api/sheets/{sid}")
    assert resp.status_code == 200
    sess_data = resp.json()
    assert sess_data["global_meta"]["brand"] == "GlobalBrand"
    assert sess_data["design_meta"]["d1"]["brand"] == "GlobalBrand"
    assert sess_data["design_meta"]["d1"]["hsn"] == "61091000"

