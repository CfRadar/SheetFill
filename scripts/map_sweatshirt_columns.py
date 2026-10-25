import xlrd
from openpyxl.utils import get_column_letter

sweat_path = r'd:\ALL_USER_DATA\SheetFillingAutomation\C_sweatshirt_810b72b2a99846d1_2510-1554FK_REQXKQO6QIEHF.xls'
wb_sweat = xlrd.open_workbook(sweat_path)
sh_sweat = wb_sweat.sheet_by_name('sweatshirt')

tshirt_path = r'templates/Flipkart_TShirt_Template.xls'
wb_tshirt = xlrd.open_workbook(tshirt_path)
sh_tshirt = wb_tshirt.sheet_by_name('t_shirt')

sweat_cols = {get_column_letter(c+1): str(sh_sweat.cell_value(0, c)).strip() for c in range(sh_sweat.ncols)}
tshirt_cols = {get_column_letter(c+1): str(sh_tshirt.cell_value(0, c)).strip() for c in range(sh_tshirt.ncols)}

# Map of header names to T-shirt col
tshirt_by_name = {v.lower(): k for k, v in tshirt_cols.items() if v}

print(f"{'Col':4s} | {'Sweatshirt Header':35s} | {'T-Shirt Col':11s} | {'Notes'}")
print("-" * 80)
for c in range(sh_sweat.ncols):
    col = get_column_letter(c + 1)
    s_hdr = sweat_cols[col]
    t_col = tshirt_by_name.get(s_hdr.lower(), "-")
    same_col = tshirt_cols.get(col, "")
    
    note = ""
    if same_col == s_hdr:
        note = "Exact same column"
    elif t_col != "-":
        note = f"Moved from {t_col}"
    else:
        note = "Sweatshirt-only column"
    print(f"{col:4s} | {s_hdr:35s} | {t_col:11s} | {note}")
