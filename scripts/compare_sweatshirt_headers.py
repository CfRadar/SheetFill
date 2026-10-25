import xlrd
from openpyxl.utils import get_column_letter

sweat_path = r'd:\ALL_USER_DATA\SheetFillingAutomation\C_sweatshirt_810b72b2a99846d1_2510-1554FK_REQXKQO6QIEHF.xls'
wb_sweat = xlrd.open_workbook(sweat_path)
sh_sweat = wb_sweat.sheet_by_name('sweatshirt')

print(f"Total Rows: {sh_sweat.nrows}, Total Cols: {sh_sweat.ncols}")

print("\n--- ALL SWEATSHIRT COLUMNS (Row 1 headers) ---")
for c in range(sh_sweat.ncols):
    col_letter = get_column_letter(c + 1)
    hdr = sh_sweat.cell_value(0, c)
    type_hint = sh_sweat.cell_value(1, c) if sh_sweat.nrows > 1 else ""
    desc = sh_sweat.cell_value(3, c) if sh_sweat.nrows > 3 else ""
    print(f"{col_letter:4s} (col {c+1:2d}): {hdr}")

print("\n--- DATA ROWS IN SWEATSHIRT (Row 5+) ---")
for r in range(4, sh_sweat.nrows):
    row_data = {}
    for c in range(sh_sweat.ncols):
        val = sh_sweat.cell_value(r, c)
        if val != "" and val is not None:
            col_letter = get_column_letter(c + 1)
            hdr = sh_sweat.cell_value(0, c)
            row_data[f"{col_letter} ({hdr})"] = val
    print(f"Row {r+1}: {row_data}")
